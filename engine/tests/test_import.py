"""Transcript / chat import and the widened audio-video decoding."""
import sqlite3
import time
import zipfile

import numpy as np
import pytest
from docx import Document
from fastapi.testclient import TestClient

from minutetaker import audio, minutes as mom, transcripts as tr
from minutetaker.api import create_app
from minutetaker.db import Database
from minutetaker.export import to_docx
from minutetaker.pipeline import Engine

TEAMS_VTT = """WEBVTT

6c4f3c2a-1d2e-4b5f-9a8b-0c1d2e3f4a5b/12-0
00:00:03.140 --> 00:00:07.230
<v Panos Adamopoulos>Good morning, let's review the migration status.</v>

6c4f3c2a-1d2e-4b5f-9a8b-0c1d2e3f4a5b/13-0
00:00:08.000 --> 00:00:12.500
<v Maria K>There is a risk with the reports module.</v>

6c4f3c2a-1d2e-4b5f-9a8b-0c1d2e3f4a5b/14-0
00:00:13.000 --> 00:00:16.000
<v Panos Adamopoulos>We decided to go live on the 20th of October.</v>
"""

TEAMS_DOCX_LINES = ["Weekly sync-20261005_100000-Meeting Recording", "October 5, 2026, 10:00AM", "45m 12s", "",
                    "Panos Adamopoulos started transcription", "",
                    "Panos Adamopoulos   0:03", "Good morning everyone.", "",
                    "Maria K   0:15", "The database cutover is ready.", "There is a risk with reports.", "",
                    "Panos Adamopoulos   1:02:05", "We decided to go live on the 20th."]

WHATSAPP_ANDROID = """05/10/2026, 10:02 - Messages and calls are end-to-end encrypted. No one outside of this chat can read them.
05/10/2026, 10:02 - Panos: Good morning team
05/10/2026, 10:02 - Panos: We need to decide on the go-live
second line of the same message
05/10/2026, 10:15 - Μαρία: <Media omitted>
05/10/2026, 10:16 - Μαρία: We decided to go live Friday <This message was edited>
13/10/2026, 09:00 - Panos: I'll send the rollback plan by Monday
"""


def _speakers(p):
    return [s["speaker"] for s in p.segments]


# ------------------------------------------------------------------ parsing --
def test_teams_vtt_voice_tags():
    p = tr.parse_cues(TEAMS_VTT)
    assert p.platform == "MS Teams" and p.kind == "transcript"
    assert _speakers(p) == ["Panos Adamopoulos", "Maria K", "Panos Adamopoulos"]
    assert p.segments[0]["start"] == 3.14 and p.segments[0]["end"] == 7.23
    assert "/12-0" not in " ".join(s["text"] for s in p.segments)      # cue ids dropped


def test_zoom_vtt_and_plain_srt():
    z = tr.parse_cues("WEBVTT\n\n1\n00:00:00.000 --> 00:00:04.000\nPanos: We decided to go live.\n\n"
                      "2\n00:00:04.500 --> 00:00:09.000\nMaria K: I'll send the plan by Friday.\n")
    assert z.platform == "Zoom" and _speakers(z) == ["Panos", "Maria K"]
    assert z.segments[1]["text"] == "I'll send the plan by Friday."
    s = tr.parse_cues("1\n00:00:01,000 --> 00:00:03,000\nHello there\nsecond line\n\n"
                      "2\n00:00:04,000 --> 00:00:06,000\nNext caption\n")
    assert _speakers(s) == ["Speaker 1"] * 2 and s.segments[0]["text"] == "Hello there second line"
    assert s.platform is None


def test_teams_docx_layout():
    p = tr.parse_text("\n".join(TEAMS_DOCX_LINES))
    assert p.platform == "MS Teams"
    assert [(s["speaker"], s["start"]) for s in p.segments] == [
        ("Panos Adamopoulos", 3.0), ("Maria K", 15.0), ("Panos Adamopoulos", 3725.0)]
    assert p.segments[1]["text"] == "The database cutover is ready. There is a risk with reports."


def test_meet_transcript_with_markers_and_attendees():
    p = tr.parse_text("Weekly sync (2026-10-05 10:00 GMT+3) - Transcript\nAttendees\nPanos Adamopoulos, Maria K\n"
                      "Transcript\n00:00:00\n\nPanos Adamopoulos: Good morning.\nMaria K: The cutover is ready.\n"
                      "00:05:00\n\nPanos Adamopoulos: Great, we decided to ship.\nMeeting ended after 00:06:12")
    assert p.platform == "Google Meet" and p.attendees == ["Panos Adamopoulos", "Maria K"]
    assert [s["start"] for s in p.segments][::2] == [0.0, 300.0]
    assert len(p.segments) == 3


def test_line_formats_and_wallclock_rebase():
    zc = tr.parse_text("10:02:33 From Panos Adamopoulos to Everyone:\n\tCan everyone see my screen?\n"
                       "10:03:01 From Maria K to Everyone:\n\tYes\n")
    assert zc.platform == "Zoom" and [s["start"] for s in zc.segments] == [0.0, 28.0]
    b = tr.parse_text("[00:00:05] Panos: Hello\n[00:00:12] Maria: Quick update.\ncontinued here\n[00:01:00] Panos: Thanks")
    assert _speakers(b) == ["Panos", "Maria", "Panos"] and b.segments[1]["text"] == "Quick update. continued here"


def test_plain_notes_are_not_mistaken_for_speakers():
    p = tr.parse_text("We reviewed the plan.\n\nNote: reports module is at risk.\nDecision: go live on 20 October.\n\n"
                      "Maria will prepare the rollback plan.")
    assert _speakers(p) == ["Speaker 1"] * 3
    assert p.segments[1]["text"].startswith("Note: reports")
    starts = [s["start"] for s in p.segments]
    assert starts == sorted(starts) and starts[1] > 0                 # estimated timeline


def test_whatsapp_android():
    p = tr.parse_text(WHATSAPP_ANDROID)
    assert p.kind == "chat" and p.platform == "WhatsApp" and p.started_at == "2026-10-05T10:02:00"
    assert _speakers(p) == ["Panos", "Panos", "Μαρία", "Panos"]       # system line + media dropped
    assert p.segments[1]["text"].endswith("second line of the same message")
    assert p.segments[2]["text"] == "We decided to go live Friday"    # "edited" marker stripped
    assert p.segments[1]["start"] > p.segments[0]["start"]            # same-minute messages stay ordered
    assert p.segments[3]["start"] == 8 * 86400 - 62 * 60              # 13/10 09:00 vs 05/10 10:02


def test_whatsapp_ios_us_dates_and_greek_ampm():
    p = tr.parse_text("[10/5/26, 10:02:33 AM] Panos: Morning\n[10/5/26, 10:05:10 AM] Maria: ‎image omitted\n"
                      "[10/13/26, 1:15:00 PM] Maria: Let's ship\n")
    assert p.started_at == "2026-10-05T10:02:33" and _speakers(p) == ["Panos", "Maria"]
    assert p.segments[1]["start"] == 8 * 86400 + 3 * 3600 + 12 * 60 + 27
    g = tr.parse_text("5/10/26, 10:02 π.μ. - Πάνος: Καλημέρα\n5/10/26, 1:30 μ.μ. - Μαρία: Τέλεια\n")
    assert g.segments[1]["start"] == (13 * 60 + 30 - 10 * 60 - 2) * 60
    assert tr.detect_language(g.segments) == "el" and tr.detect_language(p.segments) == "en"


def test_file_readers(tmp_path):
    d = tmp_path / "teams.docx"
    doc = Document()
    for ln in TEAMS_DOCX_LINES:
        doc.add_paragraph(ln)
    doc.save(str(d))
    assert _speakers(tr.parse_file(d))[:2] == ["Panos Adamopoulos", "Maria K"]

    z = tmp_path / "WhatsApp Chat with Team.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("_chat.txt", WHATSAPP_ANDROID)
        zf.writestr("IMG-0001.jpg", b"\xff\xd8")
    assert tr.parse_file(z).kind == "chat"

    pdf = tmp_path / "meet.pdf"
    _make_pdf(pdf, ["Panos Adamopoulos: We decided to go live on Friday.", "Maria K: I will prepare the rollback plan."])
    assert _speakers(tr.parse_file(pdf)) == ["Panos Adamopoulos", "Maria K"]

    t = tmp_path / "greek.txt"
    t.write_bytes("Πάνος: Καλημέρα σε όλους\nΜαρία: Ξεκινάμε\n".encode("cp1253"))   # legacy Windows encoding
    assert _speakers(tr.parse_file(t)) == ["Πάνος", "Μαρία"]

    with pytest.raises(ValueError):
        tr.parse_text("   \n ")


def _make_pdf(path, lines):
    """Minimal one-page text PDF (Helvetica, ASCII) for extraction tests."""
    content = "BT /F1 11 Tf 50 800 Td 14 TL " + " ".join(f"({ln}) Tj T*" for ln in lines) + " ET"
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 5 0 R >> >> "
            "/Contents 4 0 R >>", f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
            "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(out)


# ----------------------------------------------------------------- pipeline --
def test_transcript_import_feeds_minutes_and_export(cfg, tmp_path):
    eng = Engine(cfg)
    src = tmp_path / "sync.vtt"
    src.write_text(TEAMS_VTT, "utf-8")
    m = eng.import_file(src, "", "Other", participants_hint="Panos Adamopoulos, Maria K")
    assert m["source"] == "transcript" and m["status"] == "transcribed" and m["platform"] == "MS Teams"
    assert m["title"] == "sync" and m["language"] == "en" and m["audio_path"] is None
    parts = eng.db.participants(m["id"])
    assert {p["display_name"]: p["is_self"] for p in parts} == {"Panos Adamopoulos": 1, "Maria K": 0}
    with pytest.raises(ValueError, match="no audio"):
        eng.transcribe(m["id"])
    eng.process_all(m["id"])
    mt = eng.db.get_meeting(m["id"])
    mins = eng.db.get_minutes(m["id"])["content"]
    assert mt["status"] == "ready" and any("go live" in d for d in mins["decisions"])
    assert any("risk" in i for i in mins["open_issues"])
    doc = Document(str(to_docx(tmp_path / "o.docx", mt, mins, eng.db.summaries(m["id"]), eng.db.segments(m["id"]))))
    assert "Decisions" in "\n".join(p.text for p in doc.paragraphs)
    assert eng.db.search("reports")


def test_chat_import_uses_real_dates(cfg):
    eng = Engine(cfg)
    m = eng.import_text(WHATSAPP_ANDROID, "Team chat")
    assert m["source"] == "chat" and m["platform"] == "WhatsApp" and m["started_at"] == "2026-10-05T10:02:00"
    text = mom.transcript_text(eng.db.segments(m["id"]), m)
    assert "[2026-10-13 09:00] Panos: I'll send the rollback plan by Monday" in text
    eng.process_all(m["id"])
    mins = eng.db.get_minutes(m["id"])["content"]
    assert mins["decisions"] and mins["action_items"][0]["owner"] == "Panos"


def test_legacy_db_gets_source_column(tmp_path):
    path = tmp_path / "old.db"
    con = sqlite3.connect(path)
    con.execute("CREATE TABLE meeting (id TEXT PRIMARY KEY, title TEXT NOT NULL, platform TEXT NOT NULL DEFAULT 'Other',"
                " started_at TEXT, duration_s REAL DEFAULT 0, language TEXT, status TEXT NOT NULL DEFAULT 'new',"
                " audio_path TEXT, mic_path TEXT, system_path TEXT, participants_hint TEXT DEFAULT '',"
                " agenda TEXT DEFAULT '', error TEXT, created_at TEXT NOT NULL)")
    con.execute("INSERT INTO meeting (id, title, created_at) VALUES ('m1', 'Old', '2026-01-01')")
    con.commit()
    con.close()
    db = Database(path)
    assert db.get_meeting("m1")["source"] == ""
    assert db.create_meeting("New", source="chat")["source"] == "chat"


# -------------------------------------------------------------------- media --
def _write_media(path, seconds=2.0, with_audio=True, with_video=False, sr=16000):
    av = pytest.importorskip("av")
    with av.open(str(path), "w") as c:
        vs = astream = None
        if with_video:
            vs = c.add_stream("mpeg4", rate=10)
            vs.width = vs.height = 64
            vs.pix_fmt = "yuv420p"
        if with_audio:
            astream = c.add_stream("aac", rate=sr, layout="mono")
        if vs:
            for i in range(int(seconds * 10)):
                f = av.VideoFrame.from_ndarray(np.full((64, 64, 3), i * 10 % 255, np.uint8), format="rgb24")
                for p in vs.encode(f.reformat(format="yuv420p")):
                    c.mux(p)
            for p in vs.encode(None):
                c.mux(p)
        if astream:
            x = (0.3 * np.sin(2 * np.pi * 220 * np.arange(int(seconds * sr)) / sr)).astype(np.float32)
            for i in range(0, len(x), 1024):
                f = av.AudioFrame.from_ndarray(x[None, i:i + 1024], format="fltp", layout="mono")
                f.sample_rate, f.pts = sr, i
                for p in astream.encode(f):
                    c.mux(p)
            for p in astream.encode(None):
                c.mux(p)


@pytest.mark.parametrize("name,video", [("call.m4a", False), ("call.mp4", True), ("call.mkv", True), ("call.mov", True)])
def test_video_and_compressed_audio_without_ffmpeg(cfg, tmp_path, monkeypatch, name, video):
    monkeypatch.setattr(audio, "ffmpeg_path", lambda: None)            # PyAV fallback path
    src = tmp_path / name
    _write_media(src, with_video=video)
    assert abs(len(audio.load_audio(src)) / audio.SR - 2.0) < 0.1
    m = Engine(cfg).import_file(src, "Call")
    assert m["source"] == "media" and abs(m["duration_s"] - 2.0) < 0.1


def test_video_without_audio_is_rejected_cleanly(cfg, tmp_path, monkeypatch):
    monkeypatch.setattr(audio, "ffmpeg_path", lambda: None)
    src = tmp_path / "screen.mp4"
    _write_media(src, with_audio=False, with_video=True)
    eng = Engine(cfg)
    with pytest.raises(ValueError, match="no audio track"):
        eng.import_file(src, "Screen")
    assert eng.db.list_meetings() == []                                # no orphaned meeting
    with pytest.raises(ValueError, match="Unsupported"):
        eng.import_file(tmp_path / "notes.xyz")


# ---------------------------------------------------------------------- api --
def _wait(client, jid, headers):
    for _ in range(200):
        j = client.get(f"/jobs/{jid}", headers=headers).json()
        if j["status"] in ("done", "error"):
            return j
        time.sleep(0.05)
    raise AssertionError("job timeout")


def test_api_transcript_imports(cfg):
    c = TestClient(create_app(cfg, token="t"))
    H = {"X-MT-Token": "t"}
    types = c.get("/settings", headers=H).json()["import_types"]
    assert ".vtt" in types["transcript"] and ".avi" in types["media"]

    r = c.post("/meetings/import-text", headers=H, json={"text": WHATSAPP_ANDROID, "title": "Chat", "template": "formal"})
    assert r.status_code == 200, r.text
    mid = r.json()["meeting"]["id"]
    assert _wait(c, r.json()["job"]["id"], H)["status"] == "done"
    d = c.get(f"/meetings/{mid}", headers=H).json()
    assert d["meeting"]["source"] == "chat" and d["minutes"] and len(d["participants"]) == 2
    assert c.post(f"/meetings/{mid}/transcribe", headers=H).status_code == 400
    assert c.get(f"/meetings/{mid}/export?format=docx", headers=H).content[:2] == b"PK"

    srt = "1\n00:00:01,000 --> 00:00:03,000\nPanos: We decided to ship.\n\n2\n00:00:04,000 --> 00:00:06,000\nMaria: OK\n"
    r = c.post("/meetings/import", headers=H, files={"file": ("captions.srt", srt.encode(), "application/x-subrip")},
               data={"auto_process": "false"})
    assert r.status_code == 200 and r.json()["job"] is None
    assert r.json()["meeting"]["status"] == "transcribed"

    assert c.post("/meetings/import-text", headers=H, json={"text": "  "}).status_code == 400
    assert c.post("/meetings/import", headers=H, files={"file": ("x.xyz", b"?", "application/octet-stream")}).status_code == 400


# ------------------------------------------------------------- voice notes --
VOICE_CHAT = """05/10/2026, 10:00 - Panos: Morning, sending a voice note
05/10/2026, 10:01 - Panos: PTT-20261005-WA0001.opus (file attached)
05/10/2026, 10:03 - Μαρία: PTT-20261005-WA0002.opus (συνημμένο αρχείο)
05/10/2026, 10:04 - Μαρία: IMG-20261005-WA0003.jpg (file attached)
05/10/2026, 10:05 - Μαρία: IMG-20261005-WA0004.jpg (file attached)
Whiteboard from today
05/10/2026, 10:06 - Panos: PTT-20261005-WA0009.opus (file attached)
"""


def _write_opus(path, seconds=1.5, sr=48000):
    av = pytest.importorskip("av")
    with av.open(str(path), "w", format="ogg") as c:
        s = c.add_stream("libopus", rate=sr, layout="mono")
        t = np.arange(int(seconds * sr)) / sr
        x = np.where((t > 0.3) & (t < 1.2), 0.3 * np.sin(2 * np.pi * 220 * t), 0).astype(np.float32)
        for i in range(0, len(x), 960):
            f = av.AudioFrame.from_ndarray(x[None, i:i + 960], format="flt", layout="mono")
            f.sample_rate, f.pts = sr, i
            for p in s.encode(f):
                c.mux(p)
        for p in s.encode(None):
            c.mux(p)


def _voice_zip(tmp_path, chat=VOICE_CHAT, sidecar=True):
    import json
    a, b = tmp_path / "PTT-20261005-WA0001.opus", tmp_path / "PTT-20261005-WA0002.opus"
    _write_opus(a)
    _write_opus(b)
    z = tmp_path / "WhatsApp Chat with Team.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("_chat.txt", chat)
        zf.write(a, "PTT-20261005-WA0001.opus")
        zf.write(b, "../../PTT-20261005-WA0002.opus")              # hostile member path
        if sidecar:   # mock STT reads <file>.transcript.json; WA0002 goes through real decoding + VAD
            zf.writestr("PTT-20261005-WA0001.opus.transcript.json", json.dumps({"language": "el", "segments": [
                {"start": 0, "end": 1.2, "text": "Αποφασίσαμε να μεταφέρουμε την κυκλοφορία την Παρασκευή."}]}))
        zf.writestr("IMG-20261005-WA0004.jpg", b"\xff\xd8")
    return z


def test_whatsapp_attachments_parsed():
    p = tr.parse_text(VOICE_CHAT + "[05/10/2026, 10:07:00] Maria: <attached: 00000012-AUDIO-2026-10-05-10-07-00.opus>\n")
    voice = [s for s in p.segments if s.get("media")]
    assert [s["media"] for s in voice] == ["PTT-20261005-WA0001.opus", "PTT-20261005-WA0002.opus",
                                           "PTT-20261005-WA0009.opus", "00000012-AUDIO-2026-10-05-10-07-00.opus"]
    assert all(s["text"] == tr.VOICE_TAG for s in voice)
    texts = [s["text"] for s in p.segments]
    assert "Whiteboard from today" in texts                           # photo caption kept
    assert not any("IMG-" in t for t in texts)                         # bare photo dropped
    assert tr.detect_language([s for s in p.segments if s.get("media")]) is None


def test_voice_notes_transcribed_into_minutes(cfg, tmp_path):
    eng = Engine(cfg)
    m = eng.import_file(_voice_zip(tmp_path), "Team chat")
    assert m["status"] == "imported" and m["source"] == "chat"
    d = cfg.audio_dir / m["id"]
    assert sorted(p.name for p in (d / "voice").glob("*.opus")) == ["PTT-20261005-WA0001.opus", "PTT-20261005-WA0002.opus"]
    assert not (cfg.audio_dir / "PTT-20261005-WA0002.opus").exists()  # zip-slip contained
    segs = eng.db.segments(m["id"])
    assert sum(1 for s in segs if s["media"]) == 2
    assert tr.VOICE_MISSING in [s["text"] for s in segs]               # WA0009 not in the zip

    eng.process_all(m["id"])
    mt = eng.db.get_meeting(m["id"])
    texts = {s["media"]: s["text"] for s in eng.db.segments(m["id"]) if s["media"]}
    assert texts["voice/PTT-20261005-WA0001.opus"] == f"{tr.VOICE_TAG} Αποφασίσαμε να μεταφέρουμε την κυκλοφορία την Παρασκευή."
    assert texts["voice/PTT-20261005-WA0002.opus"].startswith(f"{tr.VOICE_TAG} [speech")   # decoded real Opus
    assert mt["status"] == "ready" and not any(s["edited"] for s in eng.db.segments(m["id"]))
    mins = eng.db.get_minutes(m["id"])["content"]
    assert any("Αποφασίσαμε" in d for d in mins["decisions"])            # voice content reached the minutes
    assert eng.db.search("κυκλοφορία")

    assert eng.voice_notes(m["id"], pending_only=True) == []
    eng.process_all(m["id"])                                            # nothing pending: minutes only
    assert eng.transcribe_voice_notes(m["id"]) == 2                     # explicit re-transcribe redoes all


def test_voice_only_chat_takes_language_from_speech(cfg, tmp_path):
    chat = "05/10/2026, 10:01 - Panos: PTT-20261005-WA0001.opus (file attached)\n" \
           "05/10/2026, 10:03 - Panos: PTT-20261005-WA0002.opus (file attached)\n"
    eng = Engine(cfg)
    m = eng.import_file(_voice_zip(tmp_path, chat), "Voice only")
    assert m["language"] is None
    eng.transcribe(m["id"])
    assert eng.db.get_meeting(m["id"])["language"] == "el"


def test_api_voice_note_media_and_retranscribe(cfg, tmp_path):
    c = TestClient(create_app(cfg, token="t"))
    H = {"X-MT-Token": "t"}
    z = _voice_zip(tmp_path)
    r = c.post("/meetings/import", headers=H, files={"file": (z.name, z.read_bytes(), "application/zip")})
    assert r.status_code == 200, r.text
    mid = r.json()["meeting"]["id"]
    assert _wait(c, r.json()["job"]["id"], H)["status"] == "done"
    segs = c.get(f"/meetings/{mid}", headers=H).json()["segments"]
    note = next(s for s in segs if s["media"])
    plain = next(s for s in segs if not s["media"])
    got = c.get(f"/meetings/{mid}/segments/{note['id']}/media", headers=H)
    assert got.status_code == 200 and got.content[:4] == b"OggS" and got.headers["content-type"].startswith("audio/ogg")
    assert c.get(f"/meetings/{mid}/segments/{plain['id']}/media", headers=H).status_code == 404
    j = c.post(f"/meetings/{mid}/transcribe", headers=H)
    assert j.status_code == 200 and _wait(c, j.json()["id"], H)["status"] == "done"
